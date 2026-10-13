#!/usr/bin/env python3
"""Export a Markdown deliverable to Excel (one sheet per table) and/or Word.

Usage: export.py FILE.md [--xlsx] [--docx]   (no flag = both)
Prints only the output paths, so it costs almost no tokens to run.
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

STATUS_COLS = {"status": '"Not Run,Pass,Fail,Blocked,N/A"',
               "pass/fail": '"Pass,Fail,Blocked,N/A"'}
FONT = "Tahoma"  # renders Thai on every Windows/Mac Excel


def clean(cell):
    cell = re.sub(r"<br\s*/?>", "\n", cell.strip(), flags=re.I)
    cell = re.sub(r"\*\*(.+?)\*\*", r"\1", cell)
    cell = re.sub(r"`(.+?)`", r"\1", cell)
    return cell


def split_row(line):
    line = line.strip().strip("|")
    return [clean(c) for c in re.split(r"(?<!\\)\|", line)]


def parse_tables(text):
    """Yield (title, rows) for every Markdown pipe table, title = last heading."""
    title, rows, tables = "Sheet", [], []
    in_code = False
    for line in text.splitlines():
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        if line.startswith("#"):
            title = line.lstrip("#").strip() or title
        if line.strip().startswith("|"):
            if re.fullmatch(r"\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?", line.strip()):
                continue
            rows.append(split_row(line))
        elif rows:
            tables.append((title, rows))
            rows = []
    if rows:
        tables.append((title, rows))
    return tables


def sheet_name(title, used):
    name = re.sub(r"[\[\]:*?/\\]", "", title)[:28] or "Sheet"
    base, n = name, 2
    while name in used:
        name = f"{base[:25]}_{n}"
        n += 1
    used.add(name)
    return name


def to_xlsx(md_path, text):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    tables = parse_tables(text)
    if not tables:
        return None
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for title, rows in tables:
        ws = wb.create_sheet(sheet_name(title, used))
        width = max(len(r) for r in rows)
        for r in rows:
            ws.append(r + [""] * (width - len(r)))
        header = [str(c.value or "").strip().lower() for c in ws[1]]
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name=FONT, size=10, bold=c.row == 1,
                              color="FFFFFF" if c.row == 1 else "000000")
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = border
                if c.row == 1:
                    c.fill = PatternFill("solid", fgColor="714B67")  # Odoo purple
        for i, _ in enumerate(header, 1):
            longest = max(max((len(l) for l in str(ws.cell(r, i).value or "").split("\n")), default=0)
                          for r in range(1, ws.max_row + 1))
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
        for i, h in enumerate(header, 1):
            if h in STATUS_COLS and ws.max_row > 1:
                dv = DataValidation(type="list", formula1=STATUS_COLS[h], allow_blank=True)
                ws.add_data_validation(dv)
                col = get_column_letter(i)
                dv.add(f"{col}2:{col}{ws.max_row + 200}")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    out = os.path.splitext(md_path)[0] + ".xlsx"
    wb.save(out)
    return out


LUA_BR = 'function RawInline(el) if el.format:match("html") and el.text:match("^<br") then return pandoc.LineBreak() end end\n'


def to_docx(md_path):
    if not shutil.which("pandoc"):
        sys.exit("pandoc not found: install pandoc or use --xlsx only")
    out = os.path.splitext(md_path)[0] + ".docx"
    with tempfile.NamedTemporaryFile("w", suffix=".lua", delete=False) as f:
        f.write(LUA_BR)
        lua = f.name
    ref = os.path.join(os.path.dirname(__file__), "reference.docx")
    cmd = ["pandoc", md_path, "-f", "gfm+raw_html", "-o", out, "--lua-filter", lua,
           "--resource-path", os.path.dirname(os.path.abspath(md_path))]
    if os.path.exists(ref):
        cmd += ["--reference-doc", ref]
    try:
        subprocess.run(cmd, check=True, stderr=subprocess.DEVNULL)
    finally:
        os.unlink(lua)
    return out


def main():
    args = sys.argv[1:]
    files = [a for a in args if not a.startswith("--")]
    if len(files) != 1:
        sys.exit(__doc__)
    md = files[0]
    want_x = "--xlsx" in args or not any(a.startswith("--") for a in args)
    want_d = "--docx" in args or not any(a.startswith("--") for a in args)
    with open(md, encoding="utf-8") as f:
        text = f.read()
    if want_x:
        print(to_xlsx(md, text) or "no tables found, xlsx skipped")
    if want_d:
        print(to_docx(md))


if __name__ == "__main__":
    main()
